import os
from typing import Dict, List, Union, Optional
import tempfile

# File processing imports
try:
    import PyPDF2
    import pdfplumber
except ImportError:
    PyPDF2 = None
    pdfplumber = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from docx import Document
except ImportError:
    Document = None

class DocumentProcessor:
    """Handles text extraction from various document formats."""
    
    def __init__(self):
        self.supported_formats = {
            'application/pdf': self._extract_pdf_text,
            'application/vnd.openxmlformats-officedocument.presentationml.presentation': self._extract_pptx_text,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': self._extract_docx_text
        }
    
    def extract_text(self, file_path: str, file_type: str) -> Union[str, Dict[str, str]]:
        """
        Extract text from uploaded document.
        
        Args:
            file_path: Path to the uploaded file
            file_type: MIME type of the file
            
        Returns:
            Extracted text content (string or dict with sections)
        """
        try:
            if file_type in self.supported_formats:
                return self.supported_formats[file_type](file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")
        except Exception as e:
            print(f"Error extracting text: {str(e)}")
            return None
    
    def _extract_pdf_text(self, file_path: str) -> str:
        """Extract text from PDF file using pdfplumber (preferred) or PyPDF2."""
        text_content = ""
        
        # Try pdfplumber first (better text extraction)
        if pdfplumber:
            try:
                with pdfplumber.open(file_path) as pdf:
                    for page_num, page in enumerate(pdf.pages, 1):
                        page_text = page.extract_text()
                        if page_text:
                            text_content += f"\n--- Page {page_num} ---\n{page_text}\n"
                return text_content.strip()
            except Exception as e:
                print(f"pdfplumber failed: {e}")
        
        # Fallback to PyPDF2
        if PyPDF2:
            try:
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page_num, page in enumerate(pdf_reader.pages, 1):
                        page_text = page.extract_text()
                        if page_text:
                            text_content += f"\n--- Page {page_num} ---\n{page_text}\n"
                return text_content.strip()
            except Exception as e:
                print(f"PyPDF2 failed: {e}")
        
        raise ImportError("PDF processing requires 'pdfplumber' or 'PyPDF2' packages")
    
    def _extract_pptx_text(self, file_path: str) -> Dict[str, str]:
        """Extract text from PowerPoint presentation."""
        if not Presentation:
            raise ImportError("PowerPoint processing requires 'python-pptx' package")
        
        try:
            presentation = Presentation(file_path)
            slides_content = {}
            
            for slide_num, slide in enumerate(presentation.slides, 1):
                slide_text = ""
                
                # Extract text from shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_text += shape.text + "\n"
                
                # Extract text from notes
                if slide.notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text
                    if notes_text.strip():
                        slide_text += f"\nNotes: {notes_text}\n"
                
                if slide_text.strip():
                    slides_content[f"Slide {slide_num}"] = slide_text.strip()
            
            return slides_content
        
        except Exception as e:
            raise Exception(f"Error processing PowerPoint file: {str(e)}")
    
    def _extract_docx_text(self, file_path: str) -> str:
        """Extract text from Word document."""
        if not Document:
            raise ImportError("Word document processing requires 'python-docx' package")
        
        try:
            doc = Document(file_path)
            text_content = ""
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_content += cell.text + " "
                    text_content += "\n"
            
            return text_content.strip()
        
        except Exception as e:
            raise Exception(f"Error processing Word document: {str(e)}")
    
    def get_document_info(self, file_path: str, file_type: str) -> Dict[str, Union[str, int]]:
        """Get basic information about the document."""
        info = {
            "file_size": os.path.getsize(file_path),
            "file_type": file_type
        }
        
        try:
            if file_type == 'application/pdf' and pdfplumber:
                with pdfplumber.open(file_path) as pdf:
                    info["page_count"] = len(pdf.pages)
            
            elif file_type == 'application/vnd.openxmlformats-officedocument.presentationml.presentation' and Presentation:
                presentation = Presentation(file_path)
                info["slide_count"] = len(presentation.slides)
            
            elif file_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' and Document:
                doc = Document(file_path)
                info["paragraph_count"] = len(doc.paragraphs)
        
        except Exception as e:
            print(f"Could not get document info: {e}")
        
        return info

